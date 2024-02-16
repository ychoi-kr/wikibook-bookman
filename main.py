import sys
from urllib.request import urlopen
from bs4 import BeautifulSoup
import re
from docx import Document

def safe_filename(filename):
    """Remove or replace characters not allowed in file names."""
    return re.sub(r'[\\/*?:"<>|! ]', "", filename)

def fetch_book_details(url):
    """Fetch book title and code download URL from the given URL."""
    try:
        # Fetch the HTML content of the page
        response = urlopen(url)
        html = response.read()
        soup = BeautifulSoup(html, 'html.parser')

        # Extract book title
        book_title = soup.select_one('#content > div:nth-child(1) > div:nth-child(2) > h1').text.strip()

        # Extract code download URL
        code_download_url = soup.select_one('#sourcecode ul li:first-of-type a')['href'].strip()

        return book_title, code_download_url
    except Exception as e:
        print(f"An error occurred: {str(e)}")
        return None, None

def modify_and_save_document(homepage_url, code_download_url, book_title):
    """Modify the .docx file to replace placeholders and save with a new name."""
    document = Document('책사용설명서.docx')
    filename_safe_book_title = safe_filename(book_title)

    for paragraph in document.paragraphs:
        if '{homepage_url}' in paragraph.text:
            paragraph.text = paragraph.text.replace('{homepage_url}', homepage_url)
            paragraph.style = document.styles['Normal']
        if '{code_download_url}' in paragraph.text:
            paragraph.text = paragraph.text.replace('{code_download_url}', code_download_url)
            paragraph.style = document.styles['Normal']

    # Assuming a function to add hyperlinks directly is not readily available in python-docx,
    # you would typically replace the text with a hyperlink manually or explore additional libraries or approaches
    # to insert hyperlinks into the document.

    new_filename = f"{filename_safe_book_title}_책사용설명서.docx"
    document.save(new_filename)
    return new_filename


homepage_url = sys.argv[1]
book_title, code_download_url = fetch_book_details(homepage_url)

if book_title and code_download_url:
    new_doc_filename = modify_and_save_document(homepage_url, code_download_url, book_title)
    print(f"Document saved as: {new_doc_filename}")
else:
    print("Failed to fetch book details or modify the document.")

